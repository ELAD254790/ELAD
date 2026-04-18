#!/usr/bin/env python3
"""
Personal Finance Tracker - Excel Generator
Creates and updates the comprehensive financial tracking workbook.
"""

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
import json
import os
from datetime import datetime, date
import calendar

# ─── COLORS ───────────────────────────────────────────────────────────────────
C_DARK_BG      = "1A1A2E"
C_HEADER_BG    = "16213E"
C_ACCENT       = "0F3460"
C_GREEN        = "27AE60"
C_RED          = "E74C3C"
C_YELLOW       = "F39C12"
C_BLUE         = "2980B9"
C_PURPLE       = "8E44AD"
C_TEAL         = "17A589"
C_ORANGE       = "E67E22"
C_LIGHT_GRAY   = "ECF0F1"
C_MID_GRAY     = "BDC3C7"
C_WHITE        = "FFFFFF"
C_GOLD         = "F1C40F"

EXPENSE_CATEGORIES = {
    "מזון וסופר":       "🛒",
    "מסעדות וקפה":      "☕",
    "תחבורה":           "🚗",
    "בילויים":          "🎮",
    "בריאות ורפואה":    "🏥",
    "ביגוד והנעלה":     "👕",
    "חשבונות קבועים":   "📋",
    "שכר דירה":         "🏠",
    "חינוך":            "📚",
    "ספורט":            "💪",
    "מתנות":            "🎁",
    "חיסכון":           "💰",
    "השקעות":           "📈",
    "אחר":              "💳",
}

def make_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def make_font(color=C_WHITE, bold=False, size=11, name="Calibri"):
    return Font(color=color, bold=bold, size=size, name=name)

def make_border(style="thin"):
    s = Side(style=style, color="404040")
    return Border(left=s, right=s, top=s, bottom=s)

def center():
    return Alignment(horizontal="center", vertical="center", wrapText=True, readingOrder=2)

def right_align():
    return Alignment(horizontal="right", vertical="center", wrapText=True, readingOrder=2)

def apply_header_style(ws, row, col, value, bg=C_HEADER_BG, fg=C_GOLD, size=12, bold=True):
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill = make_fill(bg)
    cell.font = make_font(color=fg, bold=bold, size=size)
    cell.alignment = center()
    cell.border = make_border()
    return cell

def apply_data_style(ws, row, col, value, bg=C_DARK_BG, fg=C_WHITE, bold=False, fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill = make_fill(bg)
    cell.font = make_font(color=fg, bold=bold)
    cell.alignment = center()
    cell.border = make_border()
    if fmt:
        cell.number_format = fmt
    return cell

def set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

def set_row_height(ws, row, height):
    ws.row_dimensions[row].height = height

# ─── SHEET: DASHBOARD ─────────────────────────────────────────────────────────
def create_dashboard(wb, data):
    ws = wb.create_sheet("📊 Dashboard", 0)
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_GOLD

    # Background for whole sheet
    for r in range(1, 60):
        for c in range(1, 20):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    # Title Banner
    ws.merge_cells("A1:P1")
    title = ws["A1"]
    title.value = "💎  מעקב פיננסי אישי  💎"
    title.fill = make_fill(C_ACCENT)
    title.font = Font(color=C_GOLD, bold=True, size=20, name="Calibri")
    title.alignment = center()
    set_row_height(ws, 1, 45)

    ws.merge_cells("A2:P2")
    sub = ws["A2"]
    sub.value = f"עודכן לאחרונה: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    sub.fill = make_fill(C_ACCENT)
    sub.font = make_font(color=C_MID_GRAY, size=10)
    sub.alignment = center()
    set_row_height(ws, 2, 22)

    # ── KPI Cards ────────────────────────────────────────────────────────────
    current_month = datetime.now().strftime("%B %Y")
    month_income   = data.get("month_income", 0)
    month_expenses = data.get("month_expenses", 0)
    month_savings  = month_income - month_expenses
    savings_pct    = (month_savings / month_income * 100) if month_income > 0 else 0
    health_score   = data.get("health_score", 0)

    kpis = [
        ("💵 הכנסה חודשית",  month_income,   "₪{:,.0f}", C_GREEN,  "A"),
        ("💸 הוצאות החודש",  month_expenses, "₪{:,.0f}", C_RED,    "E"),
        ("💰 חיסכון",        month_savings,  "₪{:,.0f}", C_BLUE,   "I"),
        ("📊 % חיסכון",      savings_pct,    "{:.1f}%",  C_PURPLE, "M"),
    ]

    set_row_height(ws, 4, 18)
    for label, val, fmt, color, start_col in kpis:
        sc = openpyxl.utils.column_index_from_string(start_col)
        end_col_letter = get_column_letter(sc + 2)
        ws.merge_cells(f"{start_col}4:{end_col_letter}4")
        ws.merge_cells(f"{start_col}5:{end_col_letter}5")
        ws.merge_cells(f"{start_col}6:{end_col_letter}6")

        lbl_cell = ws.cell(4, sc, label)
        lbl_cell.fill = make_fill(color)
        lbl_cell.font = make_font(color=C_WHITE, bold=True, size=10)
        lbl_cell.alignment = center()

        val_cell = ws.cell(5, sc, fmt.format(val))
        val_cell.fill = make_fill(color)
        val_cell.font = Font(color=C_WHITE, bold=True, size=18, name="Calibri")
        val_cell.alignment = center()
        set_row_height(ws, 5, 40)

        sep = ws.cell(6, sc, "")
        sep.fill = make_fill(C_DARK_BG)
        set_row_height(ws, 6, 8)

    # ── Financial Health Score ────────────────────────────────────────────────
    set_row_height(ws, 7, 18)
    ws.merge_cells("A7:P7")
    ws["A7"].value = "━━━━━━━━━━━━━━  מד הבריאות הפיננסית שלך  ━━━━━━━━━━━━━━"
    ws["A7"].fill = make_fill(C_HEADER_BG)
    ws["A7"].font = make_font(color=C_GOLD, bold=True, size=12)
    ws["A7"].alignment = center()

    # Score display
    ws.merge_cells("A8:P9")
    score_color = C_GREEN if health_score >= 70 else (C_YELLOW if health_score >= 40 else C_RED)
    score_label = "מצוין! 🏆" if health_score >= 80 else ("טוב 👍" if health_score >= 60 else ("בינוני ⚠️" if health_score >= 40 else "דורש שיפור 🔴"))
    ws["A8"].value = f"ציון: {health_score}/100  |  {score_label}"
    ws["A8"].fill = make_fill(score_color)
    ws["A8"].font = Font(color=C_WHITE, bold=True, size=22, name="Calibri")
    ws["A8"].alignment = center()
    set_row_height(ws, 8, 45)
    set_row_height(ws, 9, 10)

    # ── Score breakdown ───────────────────────────────────────────────────────
    breakdown = data.get("score_breakdown", {})
    bd_items = [
        ("חיסכון", breakdown.get("savings", 0), 30, "A"),
        ("עמידה בתקציב", breakdown.get("budget", 0), 25, "E"),
        ("פיזור הוצאות", breakdown.get("diversity", 0), 20, "I"),
        ("קרן חירום", breakdown.get("emergency", 0), 15, "M"),
    ]
    set_row_height(ws, 10, 18)
    ws.merge_cells("A10:P10")
    ws["A10"].value = "פירוט הציון:"
    ws["A10"].fill = make_fill(C_HEADER_BG)
    ws["A10"].font = make_font(color=C_MID_GRAY, size=10)
    ws["A10"].alignment = center()

    for label, score, max_score, start_col in bd_items:
        sc = openpyxl.utils.column_index_from_string(start_col)
        ec = get_column_letter(sc + 2)
        ws.merge_cells(f"{start_col}11:{ec}11")
        ws.merge_cells(f"{start_col}12:{ec}12")
        c1 = ws.cell(11, sc, f"{label}")
        c1.fill = make_fill(C_HEADER_BG)
        c1.font = make_font(color=C_MID_GRAY, size=9)
        c1.alignment = center()
        pct = int((score / max_score) * 100) if max_score > 0 else 0
        clr = C_GREEN if pct >= 70 else (C_YELLOW if pct >= 40 else C_RED)
        c2 = ws.cell(12, sc, f"{score}/{max_score}")
        c2.fill = make_fill(clr)
        c2.font = make_font(color=C_WHITE, bold=True, size=12)
        c2.alignment = center()
    set_row_height(ws, 11, 16)
    set_row_height(ws, 12, 30)

    # ── Tips Section ──────────────────────────────────────────────────────────
    set_row_height(ws, 14, 18)
    ws.merge_cells("A14:P14")
    ws["A14"].value = "💡  טיפים פיננסיים אישיים"
    ws["A14"].fill = make_fill(C_ACCENT)
    ws["A14"].font = make_font(color=C_GOLD, bold=True, size=12)
    ws["A14"].alignment = center()

    tips = data.get("tips", [
        "הגדר תקציב חודשי לכל קטגוריה ועמוד בו",
        "שמור לפחות 3 חודשי הוצאות כקרן חירום",
        "מטרה: לחסוך לפחות 20% מההכנסה נטו",
        "הגדל חיסכון אוטומטי ביום קבלת המשכורת",
    ])
    for i, tip in enumerate(tips[:6]):
        row = 15 + i
        set_row_height(ws, row, 22)
        ws.merge_cells(f"A{row}:P{row}")
        c = ws.cell(row, 1, f"  ➤  {tip}")
        bg = "0D2137" if i % 2 == 0 else "0A1929"
        c.fill = make_fill(bg)
        c.font = make_font(color=C_LIGHT_GRAY, size=10)
        c.alignment = right_align()

    # Column widths
    for col in range(1, 17):
        set_col_width(ws, col, 12)
    ws.column_dimensions["A"].width = 18

    return ws

# ─── SHEET: MONTHLY SUMMARY ───────────────────────────────────────────────────
def create_monthly_summary(wb, data):
    ws = wb.create_sheet("📅 סיכום חודשי")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_GREEN

    for r in range(1, 100):
        for c in range(1, 15):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    # Title
    ws.merge_cells("A1:N1")
    ws["A1"].value = "📅  סיכום חודשי - הכנסות והוצאות"
    ws["A1"].fill = make_fill(C_ACCENT)
    ws["A1"].font = Font(color=C_GOLD, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = center()
    set_row_height(ws, 1, 35)

    # Headers for monthly table
    months_headers = ["חודש", "הכנסה", "הוצאות", "חיסכון", "% חיסכון", "ציון", "הערות"]
    for i, h in enumerate(months_headers):
        apply_header_style(ws, 3, i+1, h, bg=C_HEADER_BG, fg=C_GOLD)
    set_row_height(ws, 3, 28)

    col_widths = [15, 14, 14, 14, 12, 10, 30]
    for i, w in enumerate(col_widths):
        set_col_width(ws, i+1, w)

    # Monthly data
    monthly = data.get("monthly_data", [])
    for row_idx, m in enumerate(monthly):
        r = 4 + row_idx
        bg = "0D2137" if row_idx % 2 == 0 else "0A1929"
        savings = m.get("income", 0) - m.get("expenses", 0)
        savings_pct = (savings / m["income"] * 100) if m.get("income", 0) > 0 else 0
        score_color = C_GREEN if m.get("score", 0) >= 70 else (C_YELLOW if m.get("score", 0) >= 40 else C_RED)

        apply_data_style(ws, r, 1, m.get("month", ""), bg=bg)
        apply_data_style(ws, r, 2, m.get("income", 0), bg=bg, fmt='₪#,##0')
        apply_data_style(ws, r, 3, m.get("expenses", 0), bg=bg, fmt='₪#,##0')
        apply_data_style(ws, r, 4, savings, bg=bg, fg=(C_GREEN if savings >= 0 else C_RED), fmt='₪#,##0')
        apply_data_style(ws, r, 5, savings_pct/100, bg=bg, fmt='0.0%')
        c = apply_data_style(ws, r, 6, m.get("score", 0), bg=score_color, fg=C_WHITE, bold=True)
        apply_data_style(ws, r, 7, m.get("notes", ""), bg=bg)
        set_row_height(ws, r, 22)

    # ── Income Section ────────────────────────────────────────────────────────
    income_start_row = max(6, 4 + len(monthly) + 2)
    ws.merge_cells(f"A{income_start_row}:N{income_start_row}")
    ws[f"A{income_start_row}"].value = "💵  מקורות הכנסה"
    ws[f"A{income_start_row}"].fill = make_fill(C_GREEN)
    ws[f"A{income_start_row}"].font = Font(color=C_WHITE, bold=True, size=13, name="Calibri")
    ws[f"A{income_start_row}"].alignment = center()
    set_row_height(ws, income_start_row, 30)

    inc_headers = ["מקור הכנסה", "סכום נטו", "סכום ברוטו", "תאריך קבלה", "קבוע/משתנה", "הערות"]
    for i, h in enumerate(inc_headers):
        apply_header_style(ws, income_start_row+1, i+1, h, bg=C_HEADER_BG, fg=C_GREEN)
    set_row_height(ws, income_start_row+1, 25)

    income_sources = data.get("income_sources", [])
    for idx, inc in enumerate(income_sources):
        r = income_start_row + 2 + idx
        bg = "071A0D" if idx % 2 == 0 else "051409"
        apply_data_style(ws, r, 1, inc.get("source", ""), bg=bg)
        apply_data_style(ws, r, 2, inc.get("net", 0), bg=bg, fg=C_GREEN, bold=True, fmt='₪#,##0')
        apply_data_style(ws, r, 3, inc.get("gross", 0), bg=bg, fmt='₪#,##0')
        apply_data_style(ws, r, 4, inc.get("date", ""), bg=bg)
        apply_data_style(ws, r, 5, inc.get("type", "קבוע"), bg=bg)
        apply_data_style(ws, r, 6, inc.get("notes", ""), bg=bg)
        set_row_height(ws, r, 22)

    return ws

# ─── SHEET: DAILY EXPENSES ────────────────────────────────────────────────────
def create_daily_expenses(wb, data):
    ws = wb.create_sheet("💸 הוצאות יומיות")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_RED

    for r in range(1, 500):
        for c in range(1, 12):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    ws.merge_cells("A1:K1")
    ws["A1"].value = "💸  יומן הוצאות מפורט"
    ws["A1"].fill = make_fill(C_RED)
    ws["A1"].font = Font(color=C_WHITE, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = center()
    set_row_height(ws, 1, 35)

    headers = ["תאריך", "קטגוריה", "תיאור", "סכום", "שיטת תשלום", "האם הכרחי?", "הערות"]
    for i, h in enumerate(headers):
        apply_header_style(ws, 2, i+1, h, bg=C_HEADER_BG, fg=C_RED)
    set_row_height(ws, 2, 26)

    col_widths = [14, 20, 30, 12, 15, 15, 25]
    for i, w in enumerate(col_widths):
        set_col_width(ws, i+1, w)

    expenses = data.get("daily_expenses", [])
    for idx, exp in enumerate(expenses):
        r = 3 + idx
        bg = "200A0A" if idx % 2 == 0 else "180707"
        apply_data_style(ws, r, 1, exp.get("date", ""), bg=bg)
        apply_data_style(ws, r, 2, exp.get("category", ""), bg=bg, fg=C_ORANGE)
        apply_data_style(ws, r, 3, exp.get("description", ""), bg=bg)
        apply_data_style(ws, r, 4, exp.get("amount", 0), bg=bg, fg=C_RED, bold=True, fmt='₪#,##0')
        apply_data_style(ws, r, 5, exp.get("method", ""), bg=bg)
        is_needed = exp.get("essential", True)
        apply_data_style(ws, r, 6, "✅ כן" if is_needed else "❌ לא", bg=bg, fg=(C_GREEN if is_needed else C_YELLOW))
        apply_data_style(ws, r, 7, exp.get("notes", ""), bg=bg)
        set_row_height(ws, r, 22)

    return ws

# ─── SHEET: FIXED EXPENSES ────────────────────────────────────────────────────
def create_fixed_expenses(wb, data):
    ws = wb.create_sheet("📋 הוצאות קבועות")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_ORANGE

    for r in range(1, 80):
        for c in range(1, 12):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    ws.merge_cells("A1:J1")
    ws["A1"].value = "📋  הוצאות קבועות חודשיות"
    ws["A1"].fill = make_fill(C_ORANGE)
    ws["A1"].font = Font(color=C_WHITE, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = center()
    set_row_height(ws, 1, 35)

    headers = ["הוצאה", "קטגוריה", "סכום", "יום חיוב", "שיטת תשלום", "חיוני?", "אפשרי לחסוך?", "הערות"]
    for i, h in enumerate(headers):
        apply_header_style(ws, 2, i+1, h, bg=C_HEADER_BG, fg=C_ORANGE)
    set_row_height(ws, 2, 26)

    col_widths = [22, 18, 12, 12, 15, 10, 15, 25]
    for i, w in enumerate(col_widths):
        set_col_width(ws, i+1, w)

    fixed = data.get("fixed_expenses", [])
    total = 0
    for idx, exp in enumerate(fixed):
        r = 3 + idx
        bg = "1A1000" if idx % 2 == 0 else "130D00"
        apply_data_style(ws, r, 1, exp.get("name", ""), bg=bg, bold=True)
        apply_data_style(ws, r, 2, exp.get("category", ""), bg=bg, fg=C_ORANGE)
        apply_data_style(ws, r, 3, exp.get("amount", 0), bg=bg, fg=C_RED, bold=True, fmt='₪#,##0')
        apply_data_style(ws, r, 4, exp.get("day", ""), bg=bg)
        apply_data_style(ws, r, 5, exp.get("method", ""), bg=bg)
        apply_data_style(ws, r, 6, "✅" if exp.get("essential") else "❌", bg=bg)
        apply_data_style(ws, r, 7, "✅" if exp.get("saveable") else "—", bg=bg)
        apply_data_style(ws, r, 8, exp.get("notes", ""), bg=bg)
        total += exp.get("amount", 0)
        set_row_height(ws, r, 22)

    # Total row
    total_row = 3 + len(fixed)
    set_row_height(ws, total_row, 30)
    ws.merge_cells(f"A{total_row}:B{total_row}")
    c = ws.cell(total_row, 1, "סה\"כ הוצאות קבועות")
    c.fill = make_fill(C_ORANGE)
    c.font = Font(color=C_WHITE, bold=True, size=13, name="Calibri")
    c.alignment = center()
    c.border = make_border()
    ws.cell(total_row, 3, total).fill = make_fill(C_RED)
    ws.cell(total_row, 3).font = Font(color=C_WHITE, bold=True, size=13, name="Calibri")
    ws.cell(total_row, 3).alignment = center()
    ws.cell(total_row, 3).number_format = '₪#,##0'
    ws.cell(total_row, 3).border = make_border()

    return ws

# ─── SHEET: BUDGET PLAN ───────────────────────────────────────────────────────
def create_budget_plan(wb, data):
    ws = wb.create_sheet("🎯 תקציב")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_BLUE

    for r in range(1, 80):
        for c in range(1, 12):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    ws.merge_cells("A1:J1")
    ws["A1"].value = "🎯  תוכנית תקציב חודשית"
    ws["A1"].fill = make_fill(C_BLUE)
    ws["A1"].font = Font(color=C_WHITE, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = center()
    set_row_height(ws, 1, 35)

    headers = ["קטגוריה", "תקציב מתוכנן", "הוצאה בפועל", "הפרש", "% ניצול", "סטטוס"]
    for i, h in enumerate(headers):
        apply_header_style(ws, 2, i+1, h, bg=C_HEADER_BG, fg=C_BLUE)
    set_row_height(ws, 2, 26)

    col_widths = [22, 15, 15, 15, 12, 14]
    for i, w in enumerate(col_widths):
        set_col_width(ws, i+1, w)

    budgets = data.get("budgets", {})
    categories = list(EXPENSE_CATEGORIES.keys())

    for idx, cat in enumerate(categories):
        r = 3 + idx
        bg = "000D1A" if idx % 2 == 0 else "000B15"
        planned = budgets.get(cat, {}).get("planned", 0)
        actual = budgets.get(cat, {}).get("actual", 0)
        diff = planned - actual
        pct = (actual / planned) if planned > 0 else 0
        status = "✅ בתקציב" if pct <= 1 else ("⚠️ קרוב לגבול" if pct <= 1.1 else "🔴 חרג!")

        icon = EXPENSE_CATEGORIES.get(cat, "")
        apply_data_style(ws, r, 1, f"{icon} {cat}", bg=bg, bold=True)
        apply_data_style(ws, r, 2, planned, bg=bg, fg=C_BLUE, fmt='₪#,##0')
        apply_data_style(ws, r, 3, actual, bg=bg, fg=(C_RED if actual > planned else C_GREEN), fmt='₪#,##0')
        apply_data_style(ws, r, 4, diff, bg=bg, fg=(C_GREEN if diff >= 0 else C_RED), bold=True, fmt='₪#,##0')
        apply_data_style(ws, r, 5, pct, bg=bg, fmt='0%')
        status_bg = C_GREEN if pct <= 1 else (C_YELLOW if pct <= 1.1 else C_RED)
        apply_data_style(ws, r, 6, status, bg=status_bg, fg=C_WHITE, bold=True)
        set_row_height(ws, r, 22)

    return ws

# ─── SHEET: CATEGORY ANALYSIS ─────────────────────────────────────────────────
def create_analysis(wb, data):
    ws = wb.create_sheet("📈 ניתוח")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_PURPLE

    for r in range(1, 80):
        for c in range(1, 15):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    ws.merge_cells("A1:N1")
    ws["A1"].value = "📈  ניתוח הוצאות לפי קטגוריות"
    ws["A1"].fill = make_fill(C_PURPLE)
    ws["A1"].font = Font(color=C_WHITE, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = center()
    set_row_height(ws, 1, 35)

    headers = ["קטגוריה", "סה\"כ החודש", "% מסה\"כ", "ממוצע יומי", "מספר עסקאות", "הכי גבוה"]
    for i, h in enumerate(headers):
        apply_header_style(ws, 2, i+1, h, bg=C_HEADER_BG, fg=C_PURPLE)
    set_row_height(ws, 2, 26)

    col_widths = [22, 15, 12, 14, 16, 14]
    for i, w in enumerate(col_widths):
        set_col_width(ws, i+1, w)

    category_data = data.get("category_totals", {})
    total_expenses = sum(v.get("total", 0) for v in category_data.values())
    days_in_month = datetime.now().day

    colors_list = [C_GREEN, C_BLUE, C_ORANGE, C_PURPLE, C_TEAL, C_RED, C_YELLOW, C_ACCENT]
    for idx, (cat, vals) in enumerate(category_data.items()):
        r = 3 + idx
        bg = "0A001A" if idx % 2 == 0 else "070014"
        total = vals.get("total", 0)
        pct = total / total_expenses if total_expenses > 0 else 0
        daily_avg = total / days_in_month if days_in_month > 0 else 0
        count = vals.get("count", 0)
        max_single = vals.get("max", 0)

        icon = EXPENSE_CATEGORIES.get(cat, "")
        apply_data_style(ws, r, 1, f"{icon} {cat}", bg=bg, bold=True)
        apply_data_style(ws, r, 2, total, bg=bg, fg=colors_list[idx % len(colors_list)], bold=True, fmt='₪#,##0')
        apply_data_style(ws, r, 3, pct, bg=bg, fmt='0.0%')
        apply_data_style(ws, r, 4, daily_avg, bg=bg, fmt='₪#,##0.0')
        apply_data_style(ws, r, 5, count, bg=bg)
        apply_data_style(ws, r, 6, max_single, bg=bg, fmt='₪#,##0')
        set_row_height(ws, r, 22)

    # Total row
    tr = 3 + len(category_data)
    set_row_height(ws, tr, 30)
    ws.merge_cells(f"A{tr}:A{tr}")
    apply_data_style(ws, tr, 1, "סה\"כ", bg=C_PURPLE, fg=C_WHITE, bold=True)
    apply_data_style(ws, tr, 2, total_expenses, bg=C_PURPLE, fg=C_WHITE, bold=True, fmt='₪#,##0')
    apply_data_style(ws, tr, 3, 1.0, bg=C_PURPLE, fg=C_WHITE, bold=True, fmt='0%')

    return ws

# ─── SHEET: SALARY SLIP LOG ───────────────────────────────────────────────────
def create_salary_log(wb, data):
    ws = wb.create_sheet("📄 תלושי שכר")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_TEAL

    for r in range(1, 100):
        for c in range(1, 16):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    ws.merge_cells("A1:O1")
    ws["A1"].value = "📄  יומן תלושי שכר"
    ws["A1"].fill = make_fill(C_TEAL)
    ws["A1"].font = Font(color=C_WHITE, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = center()
    set_row_height(ws, 1, 35)

    headers = [
        "חודש", "מעסיק", "ברוטו", "נטו", "מס הכנסה", "ביטוח לאומי",
        "בריאות", "קרן פנסיה", "קרן השתלמות", "נסיעות", "שעות נוספות", "הערות"
    ]
    for i, h in enumerate(headers):
        apply_header_style(ws, 2, i+1, h, bg=C_HEADER_BG, fg=C_TEAL)
    set_row_height(ws, 2, 26)

    col_widths = [14, 16, 12, 12, 12, 14, 12, 12, 15, 12, 14, 25]
    for i, w in enumerate(col_widths):
        set_col_width(ws, i+1, w)

    salaries = data.get("salary_slips", [])
    for idx, slip in enumerate(salaries):
        r = 3 + idx
        bg = "001A17" if idx % 2 == 0 else "001210"
        apply_data_style(ws, r, 1, slip.get("month", ""), bg=bg)
        apply_data_style(ws, r, 2, slip.get("employer", ""), bg=bg)
        apply_data_style(ws, r, 3, slip.get("gross", 0), bg=bg, fg=C_GREEN, bold=True, fmt='₪#,##0')
        apply_data_style(ws, r, 4, slip.get("net", 0), bg=bg, fg=C_TEAL, bold=True, fmt='₪#,##0')
        apply_data_style(ws, r, 5, slip.get("tax", 0), bg=bg, fg=C_RED, fmt='₪#,##0')
        apply_data_style(ws, r, 6, slip.get("national_ins", 0), bg=bg, fmt='₪#,##0')
        apply_data_style(ws, r, 7, slip.get("health", 0), bg=bg, fmt='₪#,##0')
        apply_data_style(ws, r, 8, slip.get("pension", 0), bg=bg, fg=C_BLUE, fmt='₪#,##0')
        apply_data_style(ws, r, 9, slip.get("hishtalmut", 0), bg=bg, fg=C_BLUE, fmt='₪#,##0')
        apply_data_style(ws, r, 10, slip.get("travel", 0), bg=bg, fmt='₪#,##0')
        apply_data_style(ws, r, 11, slip.get("overtime", 0), bg=bg, fmt='₪#,##0')
        apply_data_style(ws, r, 12, slip.get("notes", ""), bg=bg)
        set_row_height(ws, r, 22)

    return ws

# ─── SHEET: NEXT MONTH PLAN ───────────────────────────────────────────────────
def create_next_month_plan(wb, data):
    ws = wb.create_sheet("🗓️ תכנון חודש הבא")
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = C_TEAL

    for r in range(1, 90):
        for c in range(1, 13):
            ws.cell(r, c).fill = make_fill(C_DARK_BG)

    # Title
    ws.merge_cells("A1:L1")
    ws["A1"].value = "🗓️  תכנון פיננסי — חודש הבא"
    ws["A1"].fill = make_fill(C_TEAL)
    ws["A1"].font = Font(color=C_WHITE, bold=True, size=16, name="Calibri")
    ws["A1"].alignment = center()
    set_row_height(ws, 1, 35)

    # ── Income Planning ───────────────────────────────────────────────────────
    set_row_height(ws, 2, 8)
    ws.merge_cells("A3:L3")
    ws["A3"].value = "💵  הכנסות צפויות"
    ws["A3"].fill = make_fill(C_GREEN)
    ws["A3"].font = Font(color=C_WHITE, bold=True, size=13, name="Calibri")
    ws["A3"].alignment = center()
    set_row_height(ws, 3, 28)

    inc_hdr = ["מקור", "סכום צפוי", "תאריך", "הערות"]
    for i, h in enumerate(inc_hdr):
        apply_header_style(ws, 4, i+1, h, bg=C_HEADER_BG, fg=C_GREEN)
    set_row_height(ws, 4, 24)

    income_plan = [
        ("צבא", 1462, "1 לחודש", "קבוע"),
        ("עבודה", "?", "10 לחודש", "⚠️ משתנה — עדכן בתחילת חודש"),
        ("סה\"כ צפוי", "=B5+B6", "", ""),
    ]
    for idx, (src, amt, dt, note) in enumerate(income_plan):
        r = 5 + idx
        is_total = idx == 2
        bg = C_TEAL if is_total else ("071A0D" if idx % 2 == 0 else "051409")
        fg = C_WHITE if is_total else C_GREEN
        apply_data_style(ws, r, 1, src, bg=bg, bold=is_total)
        apply_data_style(ws, r, 2, amt, bg=bg, fg=fg, bold=is_total, fmt='₪#,##0' if isinstance(amt, (int, float)) else None)
        apply_data_style(ws, r, 3, dt, bg=bg)
        apply_data_style(ws, r, 4, note, bg=bg)
        set_row_height(ws, r, 22)

    # ── Budget Plan ────────────────────────────────────────────────────────────
    set_row_height(ws, 9, 8)
    ws.merge_cells("A10:L10")
    ws["A10"].value = "🎯  תקציב מתוכנן — לפי לקחים מהחודש הנוכחי"
    ws["A10"].fill = make_fill(C_BLUE)
    ws["A10"].font = Font(color=C_WHITE, bold=True, size=13, name="Calibri")
    ws["A10"].alignment = center()
    set_row_height(ws, 10, 28)

    bud_hdr = ["קטגוריה", "הוצאה החודש", "תקציב מומלץ", "הפרש", "הסבר"]
    for i, h in enumerate(bud_hdr):
        apply_header_style(ws, 11, i+1, h, bg=C_HEADER_BG, fg=C_BLUE)
    set_row_height(ws, 11, 24)

    col_widths = [22, 15, 15, 14, 35]
    for i, w in enumerate(col_widths):
        set_col_width(ws, i+1, w)

    budgets = data.get("budgets", {})
    budget_advice = {
        "מסעדות וקפה":    "⚠️ הוצאת 811₪! הגבל ל-200₪ — בשל, תכין אוכל",
        "מתנות":           "⚠️ 860₪ החודש! הגדר 200₪ תקרה לחודש הבא",
        "חינוך":           "שיעורי נהיגה — כמה עוד נשארו לך לבחינה?",
        "מזון וסופר":      "✅ סביר — נסה להישאר מתחת ל-350₪",
        "ספורט":           "מאמן אישי — 600₪ קבוע",
        "חשבונות קבועים":  "שקול לבטל ChatGPT (71₪) — חסוך 71₪/חודש",
        "אחר":             "⚠️ 923₪! בדוק מה היה שם ותכנן מראש",
        "חיסכון":          "💰 הפרש 1,000₪ אוטומטית ב-1 לחודש!",
        "תחבורה":          "לא הוצאת — האם יש לך נסיעות?",
        "בילויים":         "✅ 66₪ — סביר",
        "בריאות ורפואה":   "✅ 120₪ — סביר",
        "ביגוד והנעלה":    "✅ 31₪ — מצוין",
    }

    row = 12
    total_actual = 0
    total_recommended = 0
    for cat, vals in budgets.items():
        actual = vals.get("actual", 0)
        recommended = vals.get("planned", 0)
        diff = recommended - actual
        bg = "000D1A" if row % 2 == 0 else "000B15"
        icon = EXPENSE_CATEGORIES.get(cat, "📌")
        apply_data_style(ws, row, 1, f"{icon} {cat}", bg=bg, bold=True)
        apply_data_style(ws, row, 2, actual, bg=bg, fg=(C_RED if actual > recommended else C_WHITE), fmt='₪#,##0')
        apply_data_style(ws, row, 3, recommended, bg=bg, fg=C_BLUE, fmt='₪#,##0')
        diff_bg = C_GREEN if diff >= 0 else C_RED
        apply_data_style(ws, row, 4, diff, bg=diff_bg, fg=C_WHITE, bold=True, fmt='₪#,##0')
        apply_data_style(ws, row, 5, budget_advice.get(cat, ""), bg=bg, fg=C_LIGHT_GRAY)
        set_row_height(ws, row, 22)
        total_actual += actual
        total_recommended += recommended
        row += 1

    # Totals
    set_row_height(ws, row, 30)
    ws.merge_cells(f"A{row}:A{row}")
    apply_data_style(ws, row, 1, "סה\"כ", bg=C_BLUE, fg=C_WHITE, bold=True)
    apply_data_style(ws, row, 2, total_actual, bg=C_RED, fg=C_WHITE, bold=True, fmt='₪#,##0')
    apply_data_style(ws, row, 3, total_recommended, bg=C_BLUE, fg=C_WHITE, bold=True, fmt='₪#,##0')
    apply_data_style(ws, row, 4, total_recommended - total_actual, bg=(C_GREEN if total_recommended >= total_actual else C_RED), fg=C_WHITE, bold=True, fmt='₪#,##0')

    # ── Savings Goal Box ───────────────────────────────────────────────────────
    goal_row = row + 2
    ws.merge_cells(f"A{goal_row}:L{goal_row}")
    savings_goal = data.get("profile", {}).get("savings_goal_monthly", 1000)
    ws[f"A{goal_row}"].value = f"🎯  יעד חיסכון חודשי: {savings_goal:,}₪  |  הוראת קבע אוטומטית ב-1 לחודש!"
    ws[f"A{goal_row}"].fill = make_fill(C_GOLD)
    ws[f"A{goal_row}"].font = Font(color="1A1A2E", bold=True, size=13, name="Calibri")
    ws[f"A{goal_row}"].alignment = center()
    set_row_height(ws, goal_row, 35)

    return ws


# ─── MAIN BUILD FUNCTION ──────────────────────────────────────────────────────
def build_workbook(data=None, output_path=None):
    if data is None:
        data = {}
    if output_path is None:
        output_path = "/home/user/ELAD/finance/מעקב_פיננסי.xlsx"

    wb = openpyxl.Workbook()
    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    create_dashboard(wb, data)
    create_monthly_summary(wb, data)
    create_daily_expenses(wb, data)
    create_fixed_expenses(wb, data)
    create_budget_plan(wb, data)
    create_analysis(wb, data)
    create_salary_log(wb, data)
    create_next_month_plan(wb, data)

    wb.save(output_path)
    print(f"✅ קובץ נשמר: {output_path}")
    return output_path

if __name__ == "__main__":
    import sys
    data_file = sys.argv[1] if len(sys.argv) > 1 else "/home/user/ELAD/finance/data.json"
    output    = sys.argv[2] if len(sys.argv) > 2 else "/home/user/ELAD/finance/מעקב_פיננסי.xlsx"

    if os.path.exists(data_file):
        with open(data_file, "r", encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = {}

    build_workbook(data, output)
